from __future__ import annotations

import csv
import hashlib
import os
import re
from pathlib import Path
from typing import Iterable

from .config import Settings
from .models import DocumentChunk
from .pdfs import chunk_text, document_id_for, extract_pdf


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".docx",
    ".xlsx",
    ".csv",
}

_XLSX_TITLE_ALIASES = {
    "channel": "渠道基础数据 channel",
    "city_and_area": "城市和区域面积 city_and_area",
    "diversion_and_segment": "分水口和渠段 diversion_and_segment",
    "station_info_698": "测站信息表",
}


def ocr_health(
    language: str,
    tesseract_cmd: str = "",
    tessdata_dir: str = "",
) -> dict[str, object]:
    try:
        import pytesseract

        _configure_tesseract(pytesseract, tesseract_cmd, tessdata_dir)
        version = str(pytesseract.get_tesseract_version()).splitlines()[0]
        languages = set(pytesseract.get_languages(config=""))
        required = [item for item in language.split("+") if item]
        missing = [item for item in required if item not in languages]
        return {
            "available": not missing,
            "version": version,
            "languages": sorted(languages),
            "missing_languages": missing,
        }
    except Exception as exc:
        return {"available": False, "error": str(exc)}


def extract_document(
    path: Path,
    settings: Settings,
) -> tuple[dict[str, str], list[DocumentChunk]]:
    """Dispatch one supported file to a provenance-preserving extractor."""

    path = path.resolve()
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document type: {suffix or '<none>'}")
    if suffix == ".pdf":
        return extract_pdf(path, settings)
    if suffix == ".docx":
        return _extract_docx(path, settings)
    if suffix == ".xlsx":
        return _extract_xlsx(path, settings)
    if suffix == ".csv":
        return _extract_csv(path, settings)
    return _extract_image(path, settings)


def _extract_docx(
    path: Path,
    settings: Settings,
) -> tuple[dict[str, str], list[DocumentChunk]]:
    from docx import Document

    document = Document(path)
    document_id = document_id_for(path)
    title = (document.core_properties.title or "").strip() or path.stem
    author = (document.core_properties.author or "").strip()
    chunks: list[DocumentChunk] = []

    paragraphs = [
        (f"段落 {index}", paragraph.text.strip())
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    chunks.extend(
        _records_to_chunks(
            document_id,
            title,
            path,
            "docx",
            "text",
            paragraphs,
            settings,
            file_name=path.name,
        )
    )

    for table_index, table in enumerate(document.tables, start=1):
        lines = []
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text.strip()) for cell in row.cells]
            if any(cells):
                lines.append("| " + " | ".join(cells) + " |")
        text = "\n".join(lines).strip()
        if text:
            chunks.append(
                DocumentChunk(
                    chunk_id=f"{document_id}-table-{table_index}",
                    document_id=document_id,
                    title=title,
                    page=0,
                    section=f"表格 {table_index}",
                    chunk_type="table",
                    text=text,
                    source_path=str(path),
                    source_type="docx",
                    location=f"表格 {table_index}",
                    file_name=path.name,
                )
            )
    return _metadata(document_id, title, author, path, "docx"), chunks


def _extract_xlsx(
    path: Path,
    settings: Settings,
) -> tuple[dict[str, str], list[DocumentChunk]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path, read_only=True, data_only=True)
    document_id = document_id_for(path)
    title = _XLSX_TITLE_ALIASES.get(path.stem.lower(), path.stem)
    chunks: list[DocumentChunk] = []
    sheet_names: list[str] = []
    total_data_rows = 0
    all_dates: list[str] = []
    try:
        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title)
            records: list[tuple[str, str]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                if not any(value.strip() for value in values):
                    continue
                last_column = max(
                    (index for index, value in enumerate(values, start=1) if value.strip()),
                    default=1,
                )
                text = " | ".join(values[:last_column])
                location = f"{sheet.title}!A{row_index}:{get_column_letter(last_column)}{row_index}"
                records.append((location, text))
            headers = records[0][1] if records else ""
            summary_chunk = None
            if records:
                data_rows = records[1:]
                total_data_rows += len(data_rows)
                summary = (
                    f"表格汇总：文件名 {path.name}；工作表 {sheet.title}；"
                    f"表头：{headers}；非空记录 {len(records)} 行；"
                    f"数据行数 {len(data_rows)}。"
                )
                date_values = [
                    match
                    for _, row_text in data_rows
                    for match in re.findall(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", row_text)
                ]
                if date_values:
                    all_dates.extend(date_values)
                    summary += f"时间范围 {date_values[0]} 至 {date_values[-1]}。"
                summary_chunk = DocumentChunk(
                    chunk_id=f"{document_id}-summary-{sheet.title}",
                    document_id=document_id,
                    title=title,
                    page=0,
                    section=f"{sheet.title} 汇总",
                    chunk_type="table_summary",
                    text=summary,
                    source_path=str(path),
                    source_type="xlsx",
                    location=f"{sheet.title} 汇总",
                    file_name=path.name,
                    sheet_name=sheet.title,
                    headers=headers,
                )
            chunks.extend(
                _records_to_chunks(
                    document_id,
                    title,
                    path,
                    "xlsx",
                    "table",
                    records,
                    settings,
                    section=sheet.title,
                    file_name=path.name,
                    sheet_name=sheet.title,
                    headers=headers,
                )
            )
            if summary_chunk is not None:
                chunks.append(summary_chunk)
        if sheet_names:
            date_range = ""
            if all_dates:
                date_range = f"时间范围 {min(all_dates)} 至 {max(all_dates)}。"
            chunks.append(
                DocumentChunk(
                    chunk_id=f"{document_id}-document-summary",
                    document_id=document_id,
                    title=title,
                    page=0,
                    section="文档汇总",
                    chunk_type="table_summary",
                    text=(
                        f"文档汇总：文件名 {path.name}；工作表数量 {len(sheet_names)}；"
                        f"工作表：{'、'.join(sheet_names)}；"
                        f"数据行数 {total_data_rows}。{date_range}"
                    ),
                    source_path=str(path),
                    source_type="xlsx",
                    location="文档汇总",
                    file_name=path.name,
                )
            )
    finally:
        workbook.close()
    return _metadata(document_id, title, "", path, "xlsx"), chunks


def _extract_csv(
    path: Path,
    settings: Settings,
) -> tuple[dict[str, str], list[DocumentChunk]]:
    document_id = document_id_for(path)
    title = path.stem
    records: list[tuple[str, str]] = []
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                for row_index, row in enumerate(csv.reader(handle), start=1):
                    if any(cell.strip() for cell in row):
                        records.append((f"第 {row_index} 行", " | ".join(row)))
            break
        except UnicodeDecodeError as exc:
            records.clear()
            last_error = exc
    else:
        raise ValueError(f"Unable to decode CSV: {path.name}") from last_error

    chunks = _records_to_chunks(
        document_id,
        title,
        path,
        "csv",
        "table",
        records,
        settings,
        section="数据表",
        file_name=path.name,
        headers=records[0][1] if records else "",
    )
    if records:
        data_rows = records[1:]
        chunks.append(
            DocumentChunk(
                chunk_id=f"{document_id}-summary",
                document_id=document_id,
                title=title,
                page=0,
                section="数据表汇总",
                chunk_type="table_summary",
                text=(
                    f"表格汇总：文件名 {path.name}；表头：{records[0][1]}；"
                    f"非空记录 {len(records)} 行；数据行数 {len(data_rows)}。"
                ),
                source_path=str(path),
                source_type="csv",
                location="数据表汇总",
                file_name=path.name,
                headers=records[0][1],
            ),
        )
    return _metadata(document_id, title, "", path, "csv"), chunks


def _extract_image(
    path: Path,
    settings: Settings,
) -> tuple[dict[str, str], list[DocumentChunk]]:
    try:
        import pytesseract
        from PIL import Image, ImageOps
    except ImportError as exc:
        raise RuntimeError("Image OCR requires Pillow and pytesseract.") from exc

    _configure_tesseract(pytesseract, settings.tesseract_cmd, settings.tessdata_dir)
    try:
        with Image.open(path) as image:
            image = ImageOps.autocontrast(ImageOps.grayscale(image))
            try:
                text = pytesseract.image_to_string(
                    image, lang=settings.ocr_language, config="--psm 3"
                )
            except TypeError:
                text = pytesseract.image_to_string(image, lang=settings.ocr_language)
            text = text.strip()
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError("Tesseract OCR is not installed or not on PATH.") from exc
    if not text:
        raise ValueError(f"No text was recognized in image: {path.name}")

    document_id = document_id_for(path)
    title = path.stem
    chunks = [
        DocumentChunk(
            chunk_id=f"{document_id}-ocr-{index}",
            document_id=document_id,
            title=title,
            page=0,
            section="图片 OCR",
            chunk_type="image_ocr",
            text=piece,
            source_path=str(path),
            source_type=path.suffix.lower().lstrip("."),
            location="图片 OCR",
            file_name=path.name,
        )
        for index, piece in enumerate(
            chunk_text(text, settings.chunk_chars, settings.chunk_overlap)
        )
    ]
    return _metadata(document_id, title, "", path, path.suffix.lower().lstrip(".")), chunks


def _configure_tesseract(pytesseract: object, command: str, tessdata_dir: str) -> None:
    if command:
        pytesseract.pytesseract.tesseract_cmd = command
    if tessdata_dir:
        os.environ["TESSDATA_PREFIX"] = tessdata_dir


def _records_to_chunks(
    document_id: str,
    title: str,
    path: Path,
    source_type: str,
    chunk_type: str,
    records: Iterable[tuple[str, str]],
    settings: Settings,
    section: str = "正文",
    file_name: str = "",
    sheet_name: str = "",
    headers: str = "",
) -> list[DocumentChunk]:
    """Pack position-labelled records without losing their source range."""

    outputs: list[DocumentChunk] = []
    current: list[tuple[str, str]] = []
    current_size = 0

    def flush() -> None:
        nonlocal current, current_size
        if not current:
            return
        location = current[0][0] if len(current) == 1 else f"{current[0][0]} - {current[-1][0]}"
        text = "\n".join(item[1] for item in current).strip()
        outputs.append(
            DocumentChunk(
                chunk_id=(
                    f"{document_id}-{source_type}-"
                    f"{hashlib.sha1(section.encode('utf-8')).hexdigest()[:8]}-"
                    f"{len(outputs)}"
                ),
                document_id=document_id,
                title=title,
                page=0,
                section=section,
                chunk_type=chunk_type,
                text=text,
                source_path=str(path),
                source_type=source_type,
                location=location,
                    file_name=file_name or path.name,
                    sheet_name=sheet_name,
                    headers=headers,
            )
        )
        current = []
        current_size = 0

    for location, text in records:
        text = text.strip()
        if not text:
            continue
        if current and current_size + len(text) + 1 > settings.chunk_chars:
            flush()
        if len(text) > settings.chunk_chars:
            flush()
            for piece in chunk_text(text, settings.chunk_chars, settings.chunk_overlap):
                current = [(location, piece)]
                current_size = len(piece)
                flush()
            continue
        current.append((location, text))
        current_size += len(text) + 1
    flush()
    return outputs


def _metadata(
    document_id: str,
    title: str,
    author: str,
    path: Path,
    source_type: str,
) -> dict[str, str]:
    return {
        "document_id": document_id,
        "title": title,
        "author": author,
        "source_path": str(path),
        "source_type": source_type,
    }
