from pathlib import Path

from PIL import Image
from docx import Document
from openpyxl import Workbook

from hydrorag.config import Settings
from hydrorag.documents import extract_document
from hydrorag.vectorstore import LocalChromaStore


def _settings(root: Path) -> Settings:
    settings = Settings.from_env(root)
    settings.ensure_dirs()
    return settings


def test_extract_docx_and_xlsx_with_locations(tmp_path):
    settings = _settings(tmp_path)

    docx_path = tmp_path / "report.docx"
    document = Document()
    document.core_properties.title = "调度报告"
    document.add_heading("运行情况", level=1)
    document.add_paragraph("G61 闸门运行稳定。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "RMSE"
    table.cell(1, 1).text = "0.0403"
    document.save(docx_path)

    metadata, chunks = extract_document(docx_path, settings)
    assert metadata["source_type"] == "docx"
    assert any(chunk.location.startswith("段落") for chunk in chunks)
    assert any(chunk.chunk_type == "table" for chunk in chunks)

    xlsx_path = tmp_path / "metrics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果"
    sheet.append(["模型", "RMSE"])
    sheet.append(["PG-LARL", 0.0403])
    workbook.save(xlsx_path)

    metadata, chunks = extract_document(xlsx_path, settings)
    assert metadata["source_type"] == "xlsx"
    assert any(chunk.chunk_type == "table_summary" for chunk in chunks)
    assert any("数据行数 1" in chunk.text for chunk in chunks)
    assert any("结果!A1" in chunk.location for chunk in chunks)
    assert any("PG-LARL" in chunk.text for chunk in chunks)


def test_extract_csv_and_image_ocr(tmp_path, monkeypatch):
    settings = _settings(tmp_path)
    csv_path = tmp_path / "gates.csv"
    csv_path.write_text("闸门,状态\nG61,正常\n", encoding="utf-8")
    metadata, chunks = extract_document(csv_path, settings)
    assert metadata["source_type"] == "csv"
    assert chunks[0].location.startswith("第 1 行")

    image_path = tmp_path / "scan.png"
    Image.new("RGB", (20, 20), "white").save(image_path)
    import pytesseract

    monkeypatch.setattr(
        pytesseract,
        "image_to_string",
        lambda image, lang: "闸门安全约束",
    )
    metadata, chunks = extract_document(image_path, settings)
    assert metadata["source_type"] == "png"
    assert chunks[0].location == "图片 OCR"
    assert "安全约束" in chunks[0].text


def test_multiformat_chunks_keep_provenance_after_chroma_round_trip(tmp_path):
    settings = _settings(tmp_path)
    store = LocalChromaStore(tmp_path / "chroma", settings.embedding_dim)

    docx_path = tmp_path / "constraints.docx"
    document = Document()
    document.add_paragraph("G61 闸门执行安全约束 alpha-control-713。")
    document.save(docx_path)
    docx_metadata, docx_chunks = extract_document(docx_path, settings)
    store.upsert(docx_chunks)

    xlsx_path = tmp_path / "metrics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "验证结果"
    sheet.append(["模型", "指标"])
    sheet.append(["PG-LARL", "beta-metric-941"])
    workbook.save(xlsx_path)
    xlsx_metadata, xlsx_chunks = extract_document(xlsx_path, settings)
    store.upsert(xlsx_chunks)

    docx_hits = store.query(
        "alpha-control-713",
        top_k=1,
        document_ids=[docx_metadata["document_id"]],
    )
    xlsx_hits = store.query(
        "beta-metric-941",
        top_k=1,
        document_ids=[xlsx_metadata["document_id"]],
    )

    assert docx_hits[0].source_type == "docx"
    assert docx_hits[0].location.startswith("段落")
    assert xlsx_hits[0].source_type == "xlsx"
    assert "验证结果!A" in xlsx_hits[0].location
    assert store.summary_chunks([xlsx_metadata["document_id"]])
